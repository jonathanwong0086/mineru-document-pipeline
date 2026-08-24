from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from docx import Document


def _atomic_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(text, encoding="utf-8")
    temp.replace(path)


def _write_docx(path: Path, markdown: str) -> None:
    document = Document()
    for line in markdown.splitlines():
        value = line.strip()
        if not value:
            continue
        if value.startswith("### "):
            document.add_heading(value[4:], level=3)
        elif value.startswith("## "):
            document.add_heading(value[3:], level=2)
        elif value.startswith("# "):
            document.add_heading(value[2:], level=1)
        else:
            document.add_paragraph(re.sub(r"[`*_]+", "", value))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def publish_bundle(bundle: dict[str, Any], markdown: str, output_root: str | Path) -> dict[str, Path]:
    output_root = Path(output_root)
    name = str(bundle["canonical_name"])
    md_path = output_root / "rag_markdown" / f"{name}.md"
    json_path = output_root / "database_json" / f"{name}.json"
    docx_path = output_root / "human_doc" / f"{name}.docx"
    _atomic_text(md_path, markdown)
    _atomic_text(json_path, json.dumps(bundle, ensure_ascii=False, indent=2))
    _write_docx(docx_path, markdown)
    return {"markdown": md_path, "json": json_path, "docx": docx_path}
